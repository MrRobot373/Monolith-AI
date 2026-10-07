"""Write pandas DataFrames to a nicely formatted .xlsx workbook.

    from xlsx_helpers import write_sheets
    write_sheets("report.xlsx", {"Summary": df1, "Details": df2},
                 money=["revenue"], percent=["share"], dates=["date"], totals=["revenue"],
                 currency="$")

Every sheet gets: a styled, frozen header row, an autofilter, sized columns and number formats.
`totals` adds a bold TOTAL row with SUM formulas for those columns.
"""
from __future__ import annotations

from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_BORDER = Border(top=Side(style="thin", color="1F2937"))


def write_sheets(path, sheets: dict[str, pd.DataFrame], money=(), percent=(), dates=(), integers=(), totals=(), currency="$"):
    wb = Workbook()
    wb.remove(wb.active)
    money_fmt = f'"{currency}"#,##0.00' if currency else "#,##0.00"
    for name, df in sheets.items():
        ws = wb.create_sheet(title=str(name)[:31])
        cols = list(df.columns)
        ws.append([str(c) for c in cols])
        for row in df.itertuples(index=False):
            ws.append([_cell(v) for v in row])
        n = len(df)
        for i, col in enumerate(cols, start=1):
            letter = get_column_letter(i)
            h = ws.cell(row=1, column=i)
            h.fill, h.font = HEADER_FILL, HEADER_FONT
            h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            fmt = None
            if col in money:
                fmt = money_fmt
            elif col in percent:
                fmt = "0.0%"
            elif col in dates or pd.api.types.is_datetime64_any_dtype(df[col]):
                fmt = "yyyy-mm-dd"
            elif col in integers or pd.api.types.is_integer_dtype(df[col]):
                fmt = "#,##0"
            elif pd.api.types.is_float_dtype(df[col]):
                fmt = "#,##0.00"
            if fmt:
                for r in range(2, n + 2 + (1 if totals else 0)):
                    ws.cell(row=r, column=i).number_format = fmt
            width = max([len(str(col))] + [len(_text(v, fmt)) for v in df[col].head(500)]) + 2
            ws.column_dimensions[letter].width = min(max(width, 8), 60)
        if totals and n:
            r = n + 2
            ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
            for i, col in enumerate(cols, start=1):
                if col in totals:
                    letter = get_column_letter(i)
                    c = ws.cell(row=r, column=i, value=f"=SUM({letter}2:{letter}{n + 1})")
                    c.font = Font(bold=True)
                for_col = ws.cell(row=r, column=i)
                for_col.border = TOTAL_BORDER
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(max(len(cols), 1))}{n + 1}"
    wb.save(path)
    return path


def _cell(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item"):  # numpy scalars
        return v.item()
    return v


def _text(v, fmt):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, (datetime, date, pd.Timestamp)):
        return "2026-01-01"
    if isinstance(v, (int, float)) and fmt:
        return f"{v:,.2f}" + ("  " if "$" in (fmt or "") else "")
    return str(v)


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 3:
        print("Usage: python3 xlsx_helpers.py input.csv output.xlsx")
        sys.exit(2)
    write_sheets(sys.argv[2], {"Data": pd.read_csv(sys.argv[1])})
    print(f"Wrote {sys.argv[2]}")
