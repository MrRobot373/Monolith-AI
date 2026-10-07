#!/usr/bin/env python3
"""Build a clean 16:9 PowerPoint deck from a Markdown outline with python-pptx.

Outline format:
    # Deck title            first H1 = title slide (next non-empty line = subtitle)
    ## Slide title          each H2 starts a slide
    - bullet                bullets (indent 2 spaces for a sub-bullet)
    ![](chart.png)          an image on the slide (bullets go left, image right)
    | a | b |               a table
    > note text             speaker notes
    ## Section: Results     a slide titled "Section: …" becomes a section divider

Usage: python3 outline_to_pptx.py outline.md deck.pptx [--accent 0072B2] [--font "Calibri"]
"""
import argparse
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

W, H = Inches(13.333), Inches(7.5)
DARK = RGBColor(0x11, 0x18, 0x27)
MUTED = RGBColor(0x6B, 0x72, 0x80)


def parse(md):
    deck = {"title": "", "subtitle": "", "slides": []}
    cur = None
    for raw in md.splitlines():
        line = raw.rstrip()
        s = line.strip()
        if not s:
            continue
        if s.startswith("# ") and not deck["title"]:
            deck["title"] = s[2:].strip()
            continue
        if deck["title"] and not deck["slides"] and cur is None and not s.startswith(("#", "-", "*", "|", ">", "!")) and not deck["subtitle"]:
            deck["subtitle"] = s
            continue
        if s.startswith("## "):
            cur = {"title": s[3:].strip(), "bullets": [], "image": None, "table": [], "notes": []}
            deck["slides"].append(cur)
            continue
        if cur is None:
            continue
        if m := re.match(r"^(\s*)[-*+]\s+(.*)", line):
            cur["bullets"].append((1 if len(m.group(1)) >= 2 else 0, m.group(2)))
        elif m := re.match(r"!\[[^\]]*\]\(([^)]+)\)", s):
            cur["image"] = m.group(1)
        elif s.startswith("|"):
            if not re.fullmatch(r"\|?[\s:|-]+\|?", s):
                cur["table"].append([c.strip() for c in s.strip("|").split("|")])
        elif s.startswith(">"):
            cur["notes"].append(s.lstrip("> "))
        else:
            cur["bullets"].append((0, s))
    return deck


def text_box(slide, x, y, w, h, text, size, bold=False, color=DARK, font="Calibri", align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, font
    return tb


def clean(t):
    return re.sub(r"\*\*([^*]+)\*\*|\*([^*]+)\*|`([^`]+)`", lambda m: next(g for g in m.groups() if g), t)


def build(deck, out, accent="0072B2", font="Calibri"):
    acc = RGBColor.from_string(accent)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    def bar(slide):
        shp = slide.shapes.add_shape(1, 0, 0, Inches(0.18), H)
        shp.fill.solid()
        shp.fill.fore_color.rgb = acc
        shp.line.fill.background()

    s = prs.slides.add_slide(blank)
    bar(s)
    text_box(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.4), deck["title"] or "Presentation", 40, True, font=font)
    if deck["subtitle"]:
        text_box(s, Inches(0.9), Inches(3.9), Inches(11.5), Inches(0.8), deck["subtitle"], 20, color=MUTED, font=font)

    for n, sl in enumerate(deck["slides"], start=2):
        s = prs.slides.add_slide(blank)
        bar(s)
        if sl["title"].lower().startswith("section:"):
            text_box(s, Inches(0.9), Inches(3.0), Inches(11.5), Inches(1.2), sl["title"].split(":", 1)[1].strip(), 36, True, color=acc, font=font)
        else:
            text_box(s, Inches(0.9), Inches(0.5), Inches(11.6), Inches(1.0), clean(sl["title"]), 30, True, font=font)
            left, top = Inches(0.9), Inches(1.7)
            body_w = Inches(6.0) if sl["image"] or sl["table"] else Inches(11.6)
            if sl["bullets"]:
                tb = s.shapes.add_textbox(left, top, body_w, Inches(5.2))
                tf = tb.text_frame
                tf.word_wrap = True
                for i, (lvl, txt) in enumerate(sl["bullets"]):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.level = lvl
                    p.space_after = Pt(10)
                    r = p.add_run()
                    r.text = ("•  " if lvl == 0 else "–  ") + clean(txt)
                    r.font.size = Pt(22 if lvl == 0 else 18)
                    r.font.color.rgb = DARK if lvl == 0 else MUTED
                    r.font.name = font
            x = left + body_w + Inches(0.4) if sl["bullets"] else left
            avail_w = Inches(12.5) - x
            if sl["image"]:
                pic = s.shapes.add_picture(sl["image"], x, top)
                scale = min(avail_w / pic.width, Inches(5.2) / pic.height, 1.0)
                pic.width, pic.height = Emu(int(pic.width * scale)), Emu(int(pic.height * scale))
            elif sl["table"]:
                rows, cols = len(sl["table"]), max(len(r) for r in sl["table"])
                tbl = s.shapes.add_table(rows, cols, x, top, avail_w, Inches(0.45) * rows).table
                for i, row in enumerate(sl["table"]):
                    for j in range(cols):
                        c = tbl.cell(i, j)
                        c.text = clean(row[j]) if j < len(row) else ""
                        c.fill.solid()
                        c.fill.fore_color.rgb = acc if i == 0 else (RGBColor(0xF3, 0xF4, 0xF6) if i % 2 else RGBColor(0xFF, 0xFF, 0xFF))
                        for p in c.text_frame.paragraphs:
                            for r in p.runs:
                                r.font.size, r.font.name = Pt(14), font
                                r.font.bold = i == 0
                                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else DARK
        text_box(s, Inches(12.3), Inches(6.95), Inches(0.8), Inches(0.4), str(n), 11, color=MUTED, font=font, align=PP_ALIGN.RIGHT)
        if sl["notes"]:
            s.notes_slide.notes_text_frame.text = "\n".join(sl["notes"])
    prs.save(out)
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("outline")
    ap.add_argument("output")
    ap.add_argument("--accent", default="0072B2")
    ap.add_argument("--font", default="Calibri")
    a = ap.parse_args()
    deck = parse(open(a.outline, encoding="utf-8").read())
    base = os.path.dirname(os.path.abspath(a.outline))
    for sl in deck["slides"]:
        if sl["image"] and not os.path.isabs(sl["image"]):
            sl["image"] = os.path.join(base, sl["image"])
    build(deck, a.output, a.accent, a.font)
    print(f"Wrote {a.output}")
