#!/usr/bin/env python3
"""Convert simple Markdown to a styled Word document (.docx) with python-docx.

Supports: # headings (1–3), paragraphs, **bold**, *italic*, `code`, bullet and numbered lists
(one level of nesting with 2+ spaces), tables (| a | b |), > quotes, --- page breaks, ```code blocks```.

Usage: python3 md_to_docx.py input.md output.docx [--title "Document title"] [--font Calibri] [--size 11]
"""
import argparse
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*"):
            par.add_run(part[1:-1]).italic = True
        elif m := re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part):
            r = par.add_run(f"{m.group(1)} ({m.group(2)})")
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x99)
        else:
            par.add_run(part)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows if not re.fullmatch(r"\|?[\s:|-]+\|?", r.strip())]
    if not cells:
        return
    width = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=width)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j in range(width):
            cell = t.cell(i, j)
            cell.text = ""
            add_inline(cell.paragraphs[0], row[j] if j < len(row) else "")
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                shade(cell, "E7EBF0")
    doc.add_paragraph()


def convert(md: str, out: str, title=None, font="Calibri", size=11):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if title:
        doc.core_properties.title = title
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code))
            r.font.name = "Consolas"
            r.font.size = Pt(size - 1)
            i += 1
            continue
        if s == "---":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif m := re.match(r"^(#{1,3})\s+(.*)", s):
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
        elif s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        elif m := re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)", line):
            nested = len(m.group(1).replace("\t", "  ")) >= 2
            numbered = m.group(2)[0].isdigit()
            style = ("List Number" if numbered else "List Bullet") + (" 2" if nested else "")
            add_inline(doc.add_paragraph(style=style), m.group(3))
        elif s.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, s.lstrip("> "))
        else:
            para = [s]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(\s*([-*+]|\d+[.)])\s|#|\||>|```|---)", lines[i + 1].strip()):
                i += 1
                para.append(lines[i].strip())
            add_inline(doc.add_paragraph(), " ".join(para))
        i += 1
    doc.save(out)
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--title")
    ap.add_argument("--font", default="Calibri")
    ap.add_argument("--size", type=int, default=11)
    a = ap.parse_args()
    convert(open(a.input, encoding="utf-8").read(), a.output, a.title, a.font, a.size)
    print(f"Wrote {a.output}")
